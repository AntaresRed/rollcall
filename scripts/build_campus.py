#!/usr/bin/env python3
"""
Build src/data/campus.json — medical help, shops, food, facilities and each
hostel's staff, for the Campus & Hostel Services screen.

    python3 scripts/build_campus.py \
        "Campus/Campus Contacts.xlsx" src/data/campus.json

Transcribed from the campus info sheet the student body keeps (Medical,
Campus Services and one tab per hostel), and a doctor's schedule.

That sheet repeats numbers the app already holds — the hostel and mess reps,
the four night messes, Mohan Da and the WH tuck shop — and its copies had
drifted: on 2026-10-27 its NH night mess number, Mohan Da's "no delivery" and
the WH tuck shop's "only call" were all out of date. So the workbook does not
carry those rows at all. The tuck shops are copied in here from tuck.json, the
one place each is written, and test/logic.mjs fails if campus.json falls
behind it. Rebuild this after rebuilding tuck.json. The night messes and reps
are not shown here at all.

  Medical        Name | Phone | Phone 2 | Note | Hours
  Doctor Hours   Day | Hours                (Monday to Sunday, the "Doctor" row's)
  Services       Group | Name | Person | Phone | Phone label | Phone 2 |
                 Phone 2 label | Hours | Delivery | Note
  NH, OH, ...    Role | Name | Phone | Phone label | Phone 2 | Phone 2 label | Note

A hostel row that says "Contact the OH Rep" gets a link to that rep in POR
details, since the instruction is only useful with a way to reach them. The
link, not a copied number, so that POR details stays the one place it lives.
"""

import json
import os
import re
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, "src", "data")

DAYS = ["Monday", "Tuesday", "Wednesday", "Thursday",
        "Friday", "Saturday", "Sunday"]

# The row the Doctor Hours sheet belongs to. The sheet's own Timings column
# for it is blank; the schedule came separately.
DOCTOR_ROW = "Doctor"

# Screen tab -> its heading, its tab label and the Services sheet groups it
# shows, in sheet order. Shops and food are one tab, every shop on campus in
# it, and the user named it (2026-10-28). The tuck shops from tuck.json lead
# it. A tab label shorter than the heading is shown with the heading above
# the list; four tabs share a 360px phone.
GROUPS = [("shops", "All Shops & Food", "All Shops & Food", ["Store", "Food"]),
          ("sports", "Sports", "Sports", ["Services"])]

# Hostel tabs, as the sheet titles them. Only the hostel's own staff: its tuck
# shop is under All Shops & Food, and the night mess and reps have screens of
# their own (the night menu, POR details) — the user dropped all three here
# (2026-10-28).
HOSTELS = [
    ("NH", "New Hostel"),
    ("OH", "Ramanujan Hostel"),
    ("LVH", "Lake View Hostel"),
    ("WH", "White Hostel"),
]

# What the Services sheet says about a tuck shop's reach that tuck.json has no
# field for.
TUCK_REACH = {"TAGORE": "WH only"}

# ---------------------------------------------------------------- corrections
#
# Declarative, and verified on every run, as in build_por.py: a workbook that
# no longer reads the wrong value fails the build rather than having a stale
# correction silently reapplied.
CORRECTIONS = [
    # (sheet, role, field, as printed, corrected, why)
    ("WH", "Electrician", "name", "Soumin", "Soumen Roy",
     "the same number as OH's Soumen Roy, and the same person (2026-10-27)"),
    # There is no LVH admin contact to give; the rep is the way in (2026-10-28).
    ("LVH", "Plumber", "note", "Contact LVH rep or LVH admin", "Contact LVH Rep",
     "the rep stands in for the admin"),
    ("LVH", "Housekeeping Incharge", "note", "Contact LVH rep or LVH admin", "Contact LVH Rep",
     "the rep stands in for the admin"),
    # OH, from the user (2026-10-28): the electrician and housekeeping are
    # reached directly; only plumbing still goes through the rep (there is no
    # permanent plumber), and its row stays as printed.
    ("OH", "Electrician", "note", "For all issues Contact the OH Rep", "",
     "Soumen Roy is called directly"),
    ("OH", "Housekeeping Incharge", "phone", "", "7044698331",
     "the sheet gave no number, only the rep"),
    ("OH", "Housekeeping Incharge", "note", "Contact the OH Rep", "",
     "reached on their own number now"),
]

MOBILE = re.compile(r"[6-9]\d{9}")
CONTACT_REP = re.compile(r"\bcontact (?:the )?(NH|OH|LVH|WH) rep\b", re.I)


def norm(v):
    return re.sub(r"\s+", " ", str(v if v is not None else "")).strip()


def die(msg):
    raise SystemExit(f"campus: {msg}")


def read(ws, want):
    """Rows as dicts keyed by the wanted headers; every header must exist."""
    rows = [[norm(c) for c in r] for r in ws.iter_rows(values_only=True)]
    rows = [r for r in rows if any(r)]
    if not rows:
        die(f"{ws.title}: empty")
    head = [h.lower() for h in rows[0]]
    idx = {}
    for w in want:
        if w.lower() not in head:
            die(f"{ws.title}: no '{w}' column (header reads {rows[0]})")
        idx[w] = head.index(w.lower())
    return [{w: (r[i] if i < len(r) else "") for w, i in idx.items()} for r in rows[1:]]


def phones(where, *pairs):
    """[(number, label), ...] -> [{number, label}], blanks dropped, every
    number a ten-digit mobile — a typo'd number is a dead call button."""
    out = []
    for number, label in pairs:
        if not number:
            continue
        digits = re.sub(r"\D", "", number)
        if not MOBILE.fullmatch(digits):
            die(f"{where}: '{number}' is not a ten-digit mobile number")
        out.append({"number": digits, "label": label or ""})
    return out


def split_phones(where, text):
    """tuck.json writes two numbers as "a / b"."""
    return phones(where, *[(p.strip(), "") for p in (text or "").split("/")])


def item(name, person="", nums=(), hours="", delivery="", note="", rep=None, menu=None):
    return {"name": name, "person": person, "phones": list(nums), "hours": hours,
            "delivery": delivery, "note": note, "rep": rep, "menu": menu}


def delivery_of(text):
    t = text.lower()
    if t in ("", "-"):
        return ""
    if t == "yes":
        return "Delivers to hostels"
    if t == "no":
        return "No delivery"
    return text     # "Only WH" and the like, kept as printed


def load(name):
    with open(os.path.join(DATA, name), encoding="utf-8") as f:
        return json.load(f)


def check_reps(por):
    """Every hostel has an 'NH Hostel Representative' in the council list —
    a "Contact the NH Rep" row links there, and must not land on nobody."""
    roles = {p.get("role") for d in por.values() for s in d.get("sections", [])
             for p in s.get("people", [])}
    for hid, _ in HOSTELS:
        if f"{hid} Hostel Representative" not in roles:
            die(f"por.json has no '{hid} Hostel Representative'")


def tuck_item(shop):
    # Whether it delivers, not what it charges: the menu screen works out the
    # fee for the student's own hostel, and a range here only competes with it.
    d = "Delivers" if shop.get("zones") or shop.get("delivery") else ""
    reach = TUCK_REACH.get(shop["tag"])
    if reach:
        d = f"{d} ({reach})" if d else reach
    return item(shop["name"], nums=split_phones(shop["name"], shop["phone"]),
                hours=shop.get("hours", ""), delivery=d, menu="tuck")


def correct(sheet, rows):
    for s, role, field, wrong, right, why in CORRECTIONS:
        if s != sheet:
            continue
        hits = [r for r in rows if r["Role"] == role]
        if len(hits) != 1:
            die(f"correction for {s} {role}: {len(hits)} rows match, expected 1")
        key = {"name": "Name", "note": "Note", "phone": "Phone"}[field]
        if hits[0][key] != wrong:
            die(f"correction for {s} {role} is stale: the sheet reads "
                f"'{hits[0][key]}', not '{wrong}' ({why}). Update CORRECTIONS.")
        hits[0][key] = right


def build(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    warnings = []
    por, tuck = load("por.json"), load("tuck.json")
    check_reps(por)

    # Doctor's week, strictly Monday to Sunday.
    week = read(wb["Doctor Hours"], ["Day", "Hours"])
    if [r["Day"] for r in week] != DAYS:
        die(f"Doctor Hours: expected {DAYS}, found {[r['Day'] for r in week]}")
    if not all(r["Hours"] for r in week):
        die("Doctor Hours: a day has no hours")

    medical = []
    for r in read(wb["Medical"], ["Name", "Phone", "Phone 2", "Note", "Hours"]):
        if not r["Name"]:
            continue
        it = item(r["Name"], nums=phones(r["Name"], (r["Phone"], ""), (r["Phone 2"], "")),
                  hours=r["Hours"], note=r["Note"])
        if r["Name"] == DOCTOR_ROW:
            it["week"] = [{"day": w["Day"], "hours": w["Hours"]} for w in week]
        medical.append(it)
    if not any("week" in m for m in medical):
        die(f"Medical: no '{DOCTOR_ROW}' row for the Doctor Hours sheet")

    services = []
    known = [g for _, _, _, gs in GROUPS for g in gs]
    for r in read(wb["Services"], ["Group", "Name", "Person", "Phone", "Phone label",
                                    "Phone 2", "Phone 2 label", "Hours", "Delivery", "Note"]):
        if not r["Name"]:
            continue
        if r["Group"] not in known:
            die(f"Services: '{r['Name']}' is in unknown group '{r['Group']}'")
        it = item(r["Name"], r["Person"],
                  phones(r["Name"], (r["Phone"], r["Phone label"]), (r["Phone 2"], r["Phone 2 label"])),
                  r["Hours"], delivery_of(r["Delivery"]), r["Note"])
        if not (it["phones"] or it["hours"] or it["note"]):
            warnings.append(f"Services: {r['Name']} has nothing to show, left out")
            continue
        services.append((r["Group"], it))

    places = []
    for gid, label, tab, groups in GROUPS:
        items = []
        if "Store" in groups:
            items += [tuck_item(s) for s in tuck["shops"]]
        items += [it for g, it in services if g in groups]
        places.append({"id": gid, "label": label, "tab": tab, "items": items})

    hostels = []
    for hid, title in HOSTELS:
        if hid not in wb.sheetnames:
            die(f"no '{hid}' sheet")
        rows = read(wb[hid], ["Role", "Name", "Phone", "Phone label",
                              "Phone 2", "Phone 2 label", "Note"])
        rows = [r for r in rows if r["Role"]]
        correct(hid, rows)
        items = []
        for r in rows:
            m = CONTACT_REP.search(r["Note"])
            rep = None
            if m:
                if m.group(1).upper() != hid:
                    die(f"{hid} {r['Role']}: note sends you to the {m.group(1)} rep")
                # Which hostel's rep, not their number: the screen links to
                # them in POR details, so a new rep needs no rebuild here.
                rep = hid
            it = item(r["Role"], r["Name"],
                      phones(f"{hid} {r['Role']}", (r["Phone"], r["Phone label"]),
                             (r["Phone 2"], r["Phone 2 label"])),
                      note=r["Note"], rep=rep)
            if not (it["phones"] or it["note"]):
                warnings.append(f"{hid}: {r['Role']} is blank on the sheet, left out")
                continue
            items.append(it)
        hostels.append({"id": hid, "name": title, "items": items})

    # One number, one contact. Two names on a number means one of them is a
    # copy-paste slip — or the same person spelled twice, as the electrician was.
    seen = {}
    every = medical + [i for p in places for i in p["items"]] + \
        [i for h in hostels for i in h["items"]]
    for it in every:
        for p in it["phones"]:
            who = f"{it['name']}|{it['person']}"
            if seen.setdefault(p["number"], who) != who:
                die(f"{p['number']} belongs to both '{seen[p['number']]}' and '{who}'")

    return {"medical": medical, "places": places, "hostels": hostels}, warnings


def main():
    if len(sys.argv) != 3:
        die("usage: build_campus.py <Campus Contacts.xlsx> <out.json>")
    out, warnings = build(sys.argv[1])
    with open(sys.argv[2], "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=2)
        f.write("\n")
    if warnings:
        print(f"{len(warnings)} thing(s) to look at:")
        for w in warnings:
            print(f"   {w}")
    print(f"   medical {len(out['medical'])}, "
          + ", ".join(f"{p['label'].lower()} {len(p['items'])}" for p in out["places"])
          + ", " + ", ".join(f"{h['id']} {len(h['items'])}" for h in out["hostels"]))


if __name__ == "__main__":
    main()
